#!/usr/bin/env python3
"""Edge-case regression tests (from roast #1). Usage: python3 edge_cases.py dist/   (exit 1 on failure)."""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "_lib"))
import openpyxl  # noqa: E402
from casekit import Checker, errors, run_cases, value  # noqa: E402


def find(path, sheet, text):
    ws = openpyxl.load_workbook(path)[sheet]
    for row in ws.iter_rows():
        for c in row:
            if c.value == text:
                return f"'{sheet}'!{c.coordinate}"
    raise KeyError(text)


def main(dist):
    src = {k: os.path.join(dist, f"Cleaning-Price-Quote-Calculator_{k}.xlsx") for k in ("UK", "US")}
    oven = find(src["UK"], "Settings", "Inside oven")
    cases = {
        "rename_extra": ("UK", {oven: "Oven deep clean"}),
        "weekly_discount": ("UK", {"QFreq": "Weekly", "QService": "Regular clean"}),
        "margin_100": ("UK", {"Margin": 1.0}),
        "not_vat_registered": ("UK", {"VATReg": "No"}),
        "crew_zero": ("UK", {"QCrew": 0}),
        "adjust_minus_150": ("UK", {"QAdj": -1.5}),
        "unknown_condition": ("UK", {"QCond": "Filthy"}),
        "blank_frequency": ("UK", {"QFreq": None}),
        "us_default": ("US", {}),
        "uk_default": ("UK", {}),
    }
    r = run_cases(src, cases)
    c = Checker()
    st = {n: value(wb, "'Quote Calculator'!E24") or "" for n, wb in r.items()}
    c.check("rename_extra warns", st["rename_extra"].startswith("⚠"), st["rename_extra"])
    c.check("rename_extra not on quote", "Inside oven" not in (value(r["rename_extra"], "ExtrasList") or ""),
            value(r["rename_extra"], "ExtrasList"))
    c.check("weekly shows below-target", st["weekly_discount"].startswith("△"), st["weekly_discount"])
    c.check("margin 100% has no errors", not errors(r["margin_100"]), errors(r["margin_100"])[:3])
    cq = r["not_vat_registered"]["Client Quote"]
    labels = [cq.cell(row=rr, column=2).value for rr in range(20, 26)]
    c.check("no 'VAT 0%' line when not registered", not any(isinstance(x, str) and "VAT 0%" in x for x in labels),
            labels)
    c.check("no 'excl. VAT' subtotal when not registered", "Subtotal" in labels and "Subtotal (excl. VAT)" not in labels,
            labels)
    big = value(r["not_vat_registered"], "'Quote Calculator'!E5") or ""
    c.check("headline says no VAT", "no VAT" in big, big)
    line = r["crew_zero"]["Client Quote"]["B12"].value or ""
    c.check("crew 0 prints 1 cleaner", "1 cleaner(s)" in line, line)
    c.check("negative adjustment warns", st["adjust_minus_150"].startswith("⚠"), st["adjust_minus_150"])
    c.check("unknown condition warns", st["unknown_condition"].startswith("⚠ Condition"), st["unknown_condition"])
    c.check("blank frequency warns", st["blank_frequency"].startswith("⚠ Frequency"), st["blank_frequency"])
    us = r["us_default"]
    hdr = [us["Settings"].cell(row=rr, column=3).value for rr in range(1, 60)]
    c.check("US rates header in sq ft", any(isinstance(h, str) and "sq ft" in h for h in hdr), "")
    c.check("US price plausible", 400 <= (value(us, "Gross") or 0) <= 460, value(us, "Gross"))
    be = r["uk_default"]["Break-even"]
    notes = [be.cell(row=rr, column=4).value for rr in range(1, 40)]
    c.check("break-even agrees with Settings by default", any(isinstance(x, str) and x.startswith("✔ Settings match")
                                                           for x in notes), notes)
    c.check("default status meets target", st["uk_default"].startswith("✔"), st["uk_default"])
    for n, wb in r.items():
        if n != "margin_100":
            c.check(f"{n}: no formula errors", not errors(wb), errors(wb)[:3])
    return c.report("pricing-calculator edge cases")


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "dist"))
