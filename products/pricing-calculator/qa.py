#!/usr/bin/env python3
"""Recalculates each built workbook in LibreOffice and checks every formula result.

Usage: python3 qa.py dist/
Fails (exit 1) on any error value or broken invariant.
"""
import glob
import os
import subprocess
import sys
import tempfile

import openpyxl

ERR_MARKERS = ("#NAME", "#VALUE", "#REF", "#DIV", "#N/A", "#NUM", "Err:", "#NULL")


def recalc(src, outdir):
    env = dict(os.environ, HOME=os.environ.get("LO_HOME", "/tmp/lohome"))
    os.makedirs(env["HOME"], exist_ok=True)
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", outdir, src],
                   check=True, capture_output=True, env=env, timeout=180)
    return os.path.join(outdir, os.path.basename(src))


def named(wb, name):
    dn = wb.defined_names[name]
    (sheet, ref), = list(dn.destinations)
    return wb[sheet][ref.replace("$", "")].value


def check(path):
    problems = []
    with tempfile.TemporaryDirectory() as td:
        out = recalc(path, td)
        wb = openpyxl.load_workbook(out, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith(ERR_MARKERS):
                        problems.append(f"{ws.title}!{c.coordinate} = {c.value}")
        v = {n: named(wb, n) for n in ["AreaM2", "ProdRate", "LabHrs", "PriceMargin", "TravelChg", "Subtotal0",
                                       "AfterMin", "FinalNet", "NetPrice", "VATAmt", "Gross", "EffRate", "BERate",
                                       "RatePerHr", "VATUsed", "CurrencyLabel", "ExtrasList", "OnSite"]}
        cq = wb["Client Quote"]
        lines = [cq.cell(row=r, column=4).value or 0 for r in range(15, 22)]
        subtotal = cq.cell(row=22, column=4).value
        total = cq.cell(row=24, column=4).value
        pl = wb["Price List"]
        price_row = [pl.cell(row=8, column=c).value for c in range(3, 10)]
        recurring = [pl.cell(row=25, column=c).value for c in range(3, 8)]
        status = wb["Quote Calculator"]["E24"].value
        tiles = [wb["Quote Log"].cell(row=5, column=c).value for c in (2, 4, 5, 6, 8, 10, 12)]
    # invariants
    if not v["FinalNet"] or v["FinalNet"] <= 0:
        problems.append("FinalNet not positive")
    if abs(sum(lines) - (subtotal or 0)) > 0.005:
        problems.append(f"quote lines {sum(lines):.2f} != subtotal {subtotal}")
    if abs((v["NetPrice"] or 0) + (v["VATAmt"] or 0) - (v["Gross"] or 0)) > 0.005:
        problems.append("net + VAT != gross")
    if abs((total or 0) - (v["Gross"] or 0)) > 0.005:
        problems.append("quote total != gross")
    if any(p in (None, "") for p in price_row):
        problems.append(f"price list gaps: {price_row}")
    if recurring and not all(recurring[i] >= recurring[i + 1] - 0.001 for i in (0,)):
        problems.append(f"weekly price above one-off: {recurring}")
    return problems, v, lines, price_row, recurring, status, tiles


if __name__ == "__main__":
    files = sorted(glob.glob(os.path.join(sys.argv[1] if len(sys.argv) > 1 else "dist", "*.xlsx")))
    bad = 0
    for fp in files:
        problems, v, lines, price_row, recurring, status, tiles = check(fp)
        print(f"== {os.path.basename(fp)}")
        print(f"   area {v['AreaM2']:.1f} m² · rate {v['ProdRate']} · labour {v['LabHrs']:.2f} h · on site "
              f"{v['OnSite']:.2f} h · net {v['NetPrice']} · {v['CurrencyLabel']} VAT {v['VATAmt']} · gross {v['Gross']}"
              f" · eff/h {v['EffRate']:.2f} · break-even/h {v['BERate']:.2f} · settings rate/h {v['RatePerHr']:.2f}")
        print(f"   quote lines {lines}")
        print(f"   price list 90m²/1500sqft row: {price_row}")
        print(f"   recurring row: {recurring}")
        print(f"   extras: {v['ExtrasList']} | status: {status}")
        print(f"   log tiles: {tiles}")
        for pr in problems:
            print("   PROBLEM:", pr)
        bad += bool(problems)
    sys.exit(1 if bad else 0)
