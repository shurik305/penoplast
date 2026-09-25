#!/usr/bin/env python3
"""Edge-case regression tests (from roast #1). Usage: python3 edge_cases.py dist/   (exit 1 on failure)."""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "_lib"))
import openpyxl  # noqa: E402
from casekit import Checker, errors, run_cases, value  # noqa: E402

Q = "'Bid Calculator'!"


def find(path, sheet, text):
    ws = openpyxl.load_workbook(path)[sheet]
    for row in ws.iter_rows():
        for c in row:
            if c.value == text:
                return f"'{sheet}'!{c.coordinate}"
    raise KeyError(text)


def proposal_text(wb):
    ws = wb["Client Proposal"]
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def main(dist):
    src = {k: os.path.join(dist, f"Commercial-Cleaning-Bid-Calculator_{k}.xlsx") for k in ("UK", "US")}
    office = find(src["UK"], "Settings", "General office / open plan")
    cases = {
        "rename_area_type": ("UK", {office: "Open-plan office"}),
        "unknown_periodic": ("UK", {f"{Q}B35": "Window polishing", f"{Q}C35": 40, f"{Q}E35": 4}),
        "blank_frequency": ("UK", {f"{Q}D16": None}),
        "size_without_type": ("UK", {f"{Q}C23": 500}),
        "margin_100": ("UK", {"Margin": 1.0}),
        "window_zero": ("UK", {"QWindow": 0}),
        "not_vat_registered": ("UK", {"VATReg": "No"}),
        "uk_default": ("UK", {}),
        "us_default": ("US", {}),
    }
    r = run_cases(src, cases)
    c = Checker()
    st = {n: value(wb, f"{Q}J26") or "" for n, wb in r.items()}
    c.check("renamed area type warns", st["rename_area_type"].startswith("⚠"), st["rename_area_type"])
    rows = [r["rename_area_type"]["Client Proposal"].cell(row=rr, column=5).value for rr in range(16, 28)]
    c.check("renamed line marked on proposal", "⚠ check" in rows, rows)
    c.check("unknown periodic warns", st["unknown_periodic"].startswith("⚠"), st["unknown_periodic"])
    c.check("unknown periodic not promised", "Window polishing" not in (value(r["unknown_periodic"], "PeriodicList") or ""),
            value(r["unknown_periodic"], "PeriodicList"))
    c.check("blank frequency warns", st["blank_frequency"].startswith("⚠"), st["blank_frequency"])
    c.check("size without type warns", st["size_without_type"].startswith("⚠"), st["size_without_type"])
    c.check("size without type not counted", value(r["size_without_type"], "TotalArea") == 418,
            value(r["size_without_type"], "TotalArea"))
    c.check("margin 100% has no errors", not errors(r["margin_100"]), errors(r["margin_100"])[:3])
    team = [t for t in proposal_text(r["window_zero"]) if t.startswith("Cleaning time:")]
    c.check("window 0: no empty 'Team of'", team and "Team of" not in team[0], team)
    texts = proposal_text(r["not_vat_registered"])
    c.check("not registered: no VAT line", not any("VAT 0%" in t for t in texts), [t for t in texts if "VAT" in t])
    c.check("not registered: plain subtotal", "Subtotal per month" in texts, "")
    c.check("default meets target", st["uk_default"].startswith("✔"), st["uk_default"])
    us = r["us_default"]
    hdr = [us["Settings"].cell(row=rr, column=3).value for rr in range(1, 80)]
    c.check("US rates header in sq ft", any(isinstance(h, str) and "sq ft" in h for h in hdr), "")
    c.check("US fee plausible", 1600 <= (value(us, "NetMonthly") or 0) <= 1800, value(us, "NetMonthly"))
    for n, wb in r.items():
        if n != "margin_100":
            c.check(f"{n}: no formula errors", not errors(wb), errors(wb)[:3])
    return c.report("commercial-bid-calculator edge cases")


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "dist"))
